# -*- coding: utf-8 -*-
"""Generación del reporte de comparación de precios en formato Excel.

Ejercicio 05: Generar un reporte en excel con los siguientes campos:
- Libro.
- Precio interno.
- Precio web.
- Diferencia.
- Fecha de extracción.

El reporte se construye a partir de los items scrapeados en Cúspide (o del
archivo CSV de alertas del Ejercicio 04) y se escribe con `openpyxl` en un
archivo `.xlsx` descargable para su análisis.
"""

import os
from datetime import datetime
from typing import Any, Dict, Iterable, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from book_manager.reportes.alertas import (
    calcular_diferencia_precio,
    leer_alertas_csv,
)

CABECERAS_REPORTE: List[str] = [
    "Libro",
    "Precio interno",
    "Precio web",
    "Diferencia",
    "Fecha de extracción",
]

NOMBRE_HOJA = "Reporte de precios"
FORMATO_MONEDA = "#,##0.00"
FORMATO_FECHA = "yyyy-mm-dd hh:mm:ss"
ANCHO_MINIMO = 12
ANCHO_MAXIMO = 60


def obtener_ruta_reporte_excel(ruta_archivo: Optional[str] = None) -> str:
    """Devuelve la ruta absoluta del archivo Excel del reporte de precios.

    Args:
        ruta_archivo: Ruta alternativa del archivo. Si es None, se usa la ruta
            por defecto en `reportes/reporte_precios.xlsx`.

    Returns:
        Ruta absoluta del archivo Excel.
    """
    if ruta_archivo:
        return os.path.abspath(ruta_archivo)
    directorio_actual = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(directorio_actual, "reporte_precios.xlsx")


def _a_float(valor: Any) -> float:
    """Convierte un valor a float devolviendo 0.0 si no es numérico."""
    try:
        if valor is None or valor == "":
            return 0.0
        return float(valor)
    except (TypeError, ValueError):
        return 0.0


def _titulo_del_libro(item: Any) -> str:
    """Obtiene el nombre del libro desde un item scrapeado."""
    titulo = str(item.get("titulo_interno", "") or "").strip()
    if not titulo:
        titulo = str(item.get("titulo_web", "") or "").strip()
    if not titulo:
        titulo = str(item.get("isbn", "") or "").strip()
    return titulo


def construir_fila_reporte(item: Any) -> Dict[str, Any]:
    """Construye la fila del reporte de precios de un item scrapeado.

    Args:
        item: Item del scraper (dict o `CuspideItem`) con los datos del
            libro interno y de la publicación web.

    Returns:
        Diccionario con los campos del reporte: libro, precio interno,
        precio web, diferencia y fecha de extracción.
    """
    precio_interno = _a_float(item.get("precio_interno", 0.0))
    precio_web = _a_float(item.get("precio", 0.0))

    diferencia = item.get("diferencia_precio")
    if diferencia is None:
        diferencia, _, _ = calcular_diferencia_precio(
            precio_interno, precio_web
        )
    else:
        diferencia = _a_float(diferencia)

    fecha = str(item.get("fecha_extraccion", "") or "").strip()
    if not fecha:
        fecha = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    return {
        "Libro": _titulo_del_libro(item),
        "Precio interno": precio_interno,
        "Precio web": precio_web,
        "Diferencia": round(diferencia, 2),
        "Fecha de extracción": fecha,
    }


def _aplicar_formato_encabezado(hoja: Worksheet) -> None:
    """Aplica el estilo visual a la fila de encabezados de la hoja."""
    relleno = PatternFill(
        start_color="FF1F4E78", end_color="FF1F4E78", fill_type="solid"
    )
    fuente = Font(bold=True, color="FFFFFFFF", size=11)
    centrado = Alignment(horizontal="center", vertical="center")

    for indice, titulo in enumerate(CABECERAS_REPORTE, start=1):
        celda = hoja.cell(row=1, column=indice, value=titulo)
        celda.fill = relleno
        celda.font = fuente
        celda.alignment = centrado

    hoja.row_dimensions[1].height = 22


def _ajustar_anchos(hoja: Worksheet) -> None:
    """Ajusta el ancho de las columnas según el contenido de las celdas."""
    for indice, titulo in enumerate(CABECERAS_REPORTE, start=1):
        letras = get_column_letter(indice)
        longitud_maxima = len(titulo)
        for celda in hoja[letras]:
            if celda.value is not None:
                longitud_maxima = max(longitud_maxima, len(str(celda.value)))
        ancho = min(max(longitud_maxima + 2, ANCHO_MINIMO), ANCHO_MAXIMO)
        hoja.column_dimensions[letras].width = ancho


def escribir_reporte_excel(
    filas: Iterable[Dict[str, Any]],
    ruta_archivo: Optional[str] = None,
) -> str:
    """Escribe las filas del reporte en un archivo Excel.

    El archivo se genera siempre, incluso sin filas, para que exista el
    reporte descargable de cada ejecución del scraper.

    Args:
        filas: Filas del reporte con los campos de `CABECERAS_REPORTE`.
        ruta_archivo: Ruta alternativa del archivo. Si es None, se usa la ruta
            por defecto en `reportes/reporte_precios.xlsx`.

    Returns:
        Ruta absoluta del archivo Excel generado.
    """
    ruta = obtener_ruta_reporte_excel(ruta_archivo)

    carpeta = os.path.dirname(ruta)
    if carpeta:
        os.makedirs(carpeta, exist_ok=True)

    libro = Workbook()
    hoja = libro.active
    hoja.title = NOMBRE_HOJA

    hoja.append(CABECERAS_REPORTE)
    for fila in filas:
        hoja.append([fila.get(campo, "") for campo in CABECERAS_REPORTE])

    _aplicar_formato_encabezado(hoja)
    _ajustar_anchos(hoja)

    for indice in range(2, hoja.max_row + 1):
        for columna in (2, 3, 4):
            celda_precio = hoja.cell(row=indice, column=columna)
            celda_precio.number_format = FORMATO_MONEDA
        celda_fecha = hoja.cell(row=indice, column=5)
        celda_fecha.number_format = FORMATO_FECHA
        celda_fecha.alignment = Alignment(horizontal="left")

    hoja.freeze_panes = "A2"
    if hoja.max_row > 1 and hoja.max_column > 0:
        hoja.auto_filter.ref = (
            f"A1:{get_column_letter(hoja.max_column)}{hoja.max_row}"
        )

    libro.save(ruta)
    return ruta


def generar_reporte_excel(
    items: Iterable[Any],
    ruta_archivo: Optional[str] = None,
) -> str:
    """Genera el reporte Excel de precios a partir de los items scrapeados.

    Args:
        items: Items scrapeados con los precios interno y web.
        ruta_archivo: Ruta alternativa del archivo Excel.

    Returns:
        Ruta absoluta del archivo Excel generado.
    """
    filas: List[Dict[str, Any]] = [
        construir_fila_reporte(item) for item in items
    ]
    return escribir_reporte_excel(filas, ruta_archivo)


def generar_reporte_desde_alertas_csv(
    ruta_alertas: Optional[str] = None,
    ruta_reporte: Optional[str] = None,
) -> str:
    """Genera el reporte Excel a partir del archivo CSV de alertas.

    Permite producir el reporte con los datos ya alertingados por el
    Ejercicio 04, sin necesidad de volver a ejecutar el scraper.

    Args:
        ruta_alertas: Ruta del CSV de alertas. Si es None, se usa la ruta por
            defecto en `reportes/alertas_precios.csv`.
        ruta_reporte: Ruta del archivo Excel a generar. Si es None, se usa la
            ruta por defecto en `reportes/reporte_precios.xlsx`.

    Returns:
        Ruta absoluta del archivo Excel generado.
    """
    alertas = leer_alertas_csv(ruta_alertas)

    filas: List[Dict[str, Any]] = []
    for alerta in alertas:
        precio_interno = _a_float(alerta.get("precio_interno", 0.0))
        precio_web = _a_float(alerta.get("precio_web", 0.0))
        diferencia = alerta.get("diferencia")
        if diferencia is None:
            diferencia, _, _ = calcular_diferencia_precio(
                precio_interno, precio_web
            )
        filas.append(
            {
                "Libro": str(
                    alerta.get("titulo_interno", "")
                    or alerta.get("titulo_web", "")
                    or alerta.get("isbn", "")
                ),
                "Precio interno": precio_interno,
                "Precio web": precio_web,
                "Diferencia": round(_a_float(diferencia), 2),
                "Fecha de extracción": str(
                    alerta.get("fecha_alerta", "")
                    or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                ),
            }
        )

    return escribir_reporte_excel(filas, ruta_reporte)


def leer_reporte_excel(
    ruta_archivo: Optional[str] = None,
) -> List[Dict[str, Any]]:
    """Lee el reporte Excel y devuelve sus filas como diccionarios.

    Args:
        ruta_archivo: Ruta alternativa del archivo. Si es None, se usa la ruta
            por defecto en `reportes/reporte_precios.xlsx`.

    Returns:
        Lista de diccionarios con los campos de cada fila del reporte.

    Raises:
        FileNotFoundError: Si el archivo de reporte no existe.
    """
    ruta = obtener_ruta_reporte_excel(ruta_archivo)
    if not os.path.exists(ruta):
        raise FileNotFoundError(
            f"El reporte Excel '{ruta}' aún no fue generado."
        )

    libro = load_workbook(ruta, data_only=True)
    hoja = libro.active

    filas: List[Dict[str, Any]] = []
    for valores in hoja.iter_rows(min_row=2, values_only=True):
        if all(valor is None for valor in valores):
            continue
        fila: Dict[str, Any] = {}
        for indice, campo in enumerate(CABECERAS_REPORTE):
            fila[campo] = valores[indice] if indice < len(valores) else None
        filas.append(fila)

    return filas


def descargar_reporte_excel(ruta_archivo: Optional[str] = None) -> str:
    """Gestiona la descarga del reporte Excel en Colab o en local.

    En Google Colab invoca files.download. En ejecución local verifica la
    existencia del archivo y devuelve su ruta absoluta para el usuario.

    Args:
        ruta_archivo: Ruta alternativa del archivo. Si es None, se usa la ruta
            por defecto en `reportes/reporte_precios.xlsx`.

    Returns:
        Mensaje con el resultado de la descarga o la ruta del archivo.

    Raises:
        FileNotFoundError: Si el archivo de reporte no existe.
    """
    ruta = obtener_ruta_reporte_excel(ruta_archivo)
    if not os.path.exists(ruta):
        raise FileNotFoundError(
            f"El reporte Excel '{ruta}' aún no fue generado."
        )

    try:
        from google.colab import files  # type: ignore

        files.download(ruta)
        return f"Descarga iniciada en Google Colab desde: {ruta}"
    except ImportError:
        return f"Archivo disponible para descarga y análisis en: {ruta}"
